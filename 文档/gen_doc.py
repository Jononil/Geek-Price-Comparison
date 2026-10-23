# -*- coding: utf-8 -*-
"""生成《极客比价系统 项目分析与设计方案》Word 文档（.docx）。"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CN_BODY = "宋体"
CN_HEAD = "黑体"
LATIN = "Times New Roman"
BLACK = RGBColor(0, 0, 0)

doc = Document()

# ---------------- 页面设置（A4） ----------------
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.54)
sec.bottom_margin = Cm(2.54)
sec.left_margin = Cm(3.0)
sec.right_margin = Cm(2.54)


def set_eastasia(obj, font):
    """给 run 或 style 设置中文字体。"""
    rpr = obj.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)


def style_font(style, cn, latin, size, bold=False, color=BLACK):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    set_eastasia(style.element, cn)


# 默认正文样式
style_font(doc.styles["Normal"], CN_BODY, LATIN, 12, False)
doc.styles["Normal"].paragraph_format.line_spacing = 1.5
doc.styles["Normal"].paragraph_format.space_after = Pt(0)

# 标题样式
style_font(doc.styles["Heading 1"], CN_HEAD, LATIN, 16, True)
doc.styles["Heading 1"].paragraph_format.space_before = Pt(18)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(10)
style_font(doc.styles["Heading 2"], CN_HEAD, LATIN, 14, True)
doc.styles["Heading 2"].paragraph_format.space_before = Pt(12)
doc.styles["Heading 2"].paragraph_format.space_after = Pt(6)
style_font(doc.styles["Heading 3"], CN_HEAD, LATIN, 12, True)
doc.styles["Heading 3"].paragraph_format.space_before = Pt(8)
doc.styles["Heading 3"].paragraph_format.space_after = Pt(4)


def h1(text):
    return doc.add_heading(text, level=1)


def h2(text):
    return doc.add_heading(text, level=2)


def h3(text):
    return doc.add_heading(text, level=3)


def p(text, indent=True, align=None):
    para = doc.add_paragraph()
    run = para.add_run(text)
    set_eastasia(run._element, CN_BODY)
    run.font.name = LATIN
    run.font.size = Pt(12)
    if indent:
        para.paragraph_format.first_line_indent = Pt(24)
    if align is not None:
        para.alignment = align
    return para


def bullet(text):
    para = doc.add_paragraph()
    run = para.add_run("• " + text)
    set_eastasia(run._element, CN_BODY)
    run.font.name = LATIN
    run.font.size = Pt(12)
    para.paragraph_format.left_indent = Pt(24)
    return para


def caption(text):
    para = doc.add_paragraph()
    run = para.add_run(text)
    set_eastasia(run._element, CN_HEAD)
    run.font.name = LATIN
    run.font.size = Pt(10.5)
    run.font.bold = True
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(3)
    return para


def shade(cell, color="D9E2F3"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def cell_text(cell, text, bold=False, size=10.5, align="left"):
    cell.text = ""
    para = cell.paragraphs[0]
    run = para.add_run(str(text))
    set_eastasia(run._element, CN_BODY)
    run.font.name = LATIN
    run.font.size = Pt(size)
    run.font.bold = bold
    if align == "center":
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(1)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        cell_text(hdr[i], htext, bold=True, align="center")
        shade(hdr[i])
    for r in rows:
        cells = t.add_row().cells
        for i, val in enumerate(r):
            cell_text(cells[i], val)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t


def page_break():
    doc.add_page_break()


def add_page_number_footer():
    footer = sec.footer
    para = footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run("第 ")
    set_eastasia(run._element, CN_BODY)
    run.font.size = Pt(9)
    r = para.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = "PAGE"
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    r._r.append(f1); r._r.append(it); r._r.append(f2)
    run2 = para.add_run(" 页")
    set_eastasia(run2._element, CN_BODY)
    run2.font.size = Pt(9)


add_page_number_footer()

# ============================================================
# 封面
# ============================================================
for _ in range(4):
    doc.add_paragraph()

para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = para.add_run("极客比价系统")
set_eastasia(run._element, CN_HEAD); run.font.name = LATIN
run.font.size = Pt(36); run.font.bold = True

para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = para.add_run("高性价比数码产品测评聚合与比价系统")
set_eastasia(run._element, CN_HEAD); run.font.name = LATIN
run.font.size = Pt(16)

doc.add_paragraph()
para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = para.add_run("项目分析与设计方案")
set_eastasia(run._element, CN_HEAD); run.font.name = LATIN
run.font.size = Pt(26); run.font.bold = True

for _ in range(5):
    doc.add_paragraph()

for line in ["第一篇    设计说明书", "第二篇    分阶段建设计划"]:
    para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(line)
    set_eastasia(run._element, CN_HEAD); run.font.name = LATIN
    run.font.size = Pt(16); run.font.bold = True

for _ in range(6):
    doc.add_paragraph()

para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = para.add_run("二〇二六年九月")
set_eastasia(run._element, CN_BODY); run.font.name = LATIN
run.font.size = Pt(16)

page_break()

# ============================================================
# 目录
# ============================================================
h1("目  录")
toc_para = doc.add_paragraph()
run = toc_para.add_run()
f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
it.text = 'TOC \\o "1-3" \\h \\z \\u'
f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
t = OxmlElement("w:t"); t.text = "（在 Word 中全选后按 F9 更新目录）"
f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
run._r.append(f1); run._r.append(it); run._r.append(f2); run._r.append(t); run._r.append(f3)

page_break()

# ============================================================
# 第一篇 设计说明书
# ============================================================
para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = para.add_run("第一篇    设计说明书")
set_eastasia(run._element, CN_HEAD); run.font.name = LATIN
run.font.size = Pt(18); run.font.bold = True
doc.add_paragraph()

# ---------- 第1章 项目概述 ----------
h1("第1章  项目概述")
h2("1.1  项目背景")
p("随着电子商务的成熟与数码产品迭代速度的加快，消费者在选购手机、笔记本电脑、耳机等数码产品时普遍面临三大难题：其一，同一产品在不同电商平台售价差异显著，人工比价耗时费力；其二，CPU、GPU、屏幕、续航等专业参数晦涩难懂，普通消费者难以据此判断产品优劣；其三，价格波动频繁，消费者无法判断「当前价格是否处于历史低位」，常常担心买在价格高点。市场上缺少一个将「专业测评、用户口碑、多平台价格、历史价格」聚合在一起，并以通俗语言输出「性价比结论」的一站式信息平台。")
p("在此背景下，本项目「极客比价」应运而生。它面向普通消费者，聚合专业测评、用户评价、多平台实时价格与历史价格，通过性价比评分算法与混合推荐算法，帮助用户快速判断「哪个值得买、现在买亏不亏」。")

h2("1.2  设计依据")
bullet("《极客比价——Claude Code 总控 Prompt.md》（项目权威需求）")
bullet("《Product-Spec.md》（产品需求文档，版本 v1.1.0）")
bullet("数据库设计规范：utf8mb4 字符集、InnoDB 存储引擎、逻辑删除（deleted 字段）约定")
bullet("《XXXX报关公司系统设计方案》（本文档章节结构与格式参考）")
bullet("软件工程、数据库系统课程教学要求（本科毕业设计）")

h2("1.3  主要业务范围")
p("本项目在统一技术框架下规划与设计三个端，覆盖从内容生产到消费决策的完整业务闭环：")
bullet("用户端：产品浏览与搜索、产品详情、多平台比价、历史价格分析、横向对比、专业测评、用户评价、收藏与浏览历史、排行榜、个性化推荐等面向消费者的功能。")
bullet("管理后台：用户、产品、分类、品牌、平台、价格、测评、评论等内容与数据管理，以及性价比评分算法参数配置与全站评分重算。")
bullet("后端 API 与数据库：为上述两端提供统一的业务逻辑、算法计算与数据持久化服务。")

h2("1.4  设计概要")
h3("1.4.1  建设目标")
p("（1）完成一个用户端、管理后台、后端 API、MySQL 数据库完整联调的真实可运行 Web 项目，满足本科毕业设计的完整度要求；")
p("（2）实现两个具有算法含量的核心能力——产品性价比评分算法与混合推荐算法，作为论文的核心技术点；")
p("（3）项目界面达到现代化科技产品网站的视觉效果，便于部署演示与后续论文撰写。")
h3("1.4.2  建设方案")
p("采用 Vue 3 + Java 全栈、前后端分离的技术路线：前端划分为用户端与管理后台两个独立项目；后端采用单体分层架构（controller / service / mapper / entity / dto / vo / config / security / exception / common）；数据库采用 MySQL 8，缓存采用 Redis（可降级）。总体方案详见第 4 章。")
h3("1.4.3  建设工期")
p("项目按「需求分析—设计—开发—联调—验收」五个阶段推进，累计覆盖 18 个开发阶段。详见第二篇第 12 章。")
h3("1.4.4  项目概算")
p("本项目为本科毕业设计，部署于个人开发机，无硬件采购成本；技术栈全部采用开源软件，软件授权成本为零。主要成本为人力（个人开发）与可选的云服务（LLM API，设计为可降级，不接入亦不影响核心功能）。")

# ---------- 第2章 建设需求 ----------
h1("第2章  建设需求")
h2("2.1  项目建设目的")
p("（1）解决数码产品选购过程中的信息不对称与比价难题，降低消费者决策成本；")
p("（2）以可解释的算法输出「性价比评分」与「个性化推荐」，避免黑盒式的购买建议；")
p("（3）完成符合本科毕业设计要求的完整全栈项目，覆盖需求、设计、开发、联调、验收全过程。")

h2("2.2  系统功能需求")
h3("2.2.1  用户端功能需求")
caption("表 2-1  用户端功能需求清单表")
table(
    ["编号", "功能模块", "功能描述", "优先级"],
    [
        ["F1", "用户系统", "注册、登录、退出、个人中心、修改密码，为个性化功能提供前置", "高"],
        ["F2", "产品浏览与搜索", "按分类/品牌浏览，关键词、价格区间、评分等多条件筛选与排序", "高"],
        ["F3", "产品详情页", "聚合参数、评分、价格、测评、评价，购买决策核心落地页", "高"],
        ["F4", "多平台比价", "京东/淘宝/拼多多等多平台价格对比，展示最低价与价差", "高"],
        ["F5", "历史价格分析", "历史价格曲线（ECharts），辅助判断「当前是否好价」", "中"],
        ["F6", "产品横向对比", "2～4 个产品参数、价格、评分并排对比", "中"],
        ["F7", "专业测评", "浏览、点赞、收藏、评论专业测评文章", "高"],
        ["F8", "用户评价", "星级评价、点赞、评论/回复，支持多次评价", "高"],
        ["F9", "收藏与浏览历史", "收藏产品/测评，自动记录浏览历史", "中"],
        ["F10", "排行榜", "性价比/热门/降价/评分/最新五类榜单", "中"],
        ["F11", "性价比评分算法", "五维加权计算 0-100 性价比评分", "高"],
        ["F12", "个性化推荐", "混合推荐 + 冷启动，为每个用户生成推荐", "低"],
        ["F13", "AI 增强功能", "AI 摘要、推荐理由、对比总结（可降级）", "低"],
    ],
    widths=[1.2, 3.2, 7.2, 1.6],
)

h3("2.2.2  管理后台功能需求")
p("管理后台面向管理员与编辑角色，提供全站内容与数据管理能力：")
bullet("用户管理：查询、禁用、删除用户；")
bullet("产品管理：新增、编辑、删除、上下架、设推荐/设热门；")
bullet("分类、品牌、平台管理：完整增删改查；")
bullet("价格管理：按产品录入各平台价格，录入后自动生成历史价格记录；")
bullet("测评、评论管理：发布/编辑/下架/删除与违规内容审核；")
bullet("算法参数配置：调整性价比评分各维度权重并重新计算全站评分；")
bullet("数据看板：统计卡片与 ECharts 图表（用户增长、分类分布、热门排行、价格趋势等）。")

h2("2.3  系统性能需求")
h3("2.3.1  系统处理能力需求")
p("系统应能支持至少 100 个并发用户正常访问，产品数据规模满足演示与论文要求（初始化数据含 132 个产品、396 条平台价格、11880 条历史价格、32 篇测评、396 条评价）。")
h3("2.3.2  系统响应速度需求")
caption("表 2-2  性能指标需求表")
table(
    ["指标", "要求"],
    [
        ["产品列表/搜索接口 P95 响应时间", "＜ 500 ms"],
        ["产品详情页接口 P95 响应时间", "＜ 800 ms"],
        ["首页首屏加载时间", "＜ 2 s"],
        ["并发访问能力", "≥ 100 并发用户"],
        ["缓存依赖", "Redis 可降级，未启动时核心功能正常"],
    ],
    widths=[8.0, 5.0],
)
h3("2.3.3  系统可靠性需求")
p("所有列表与详情页需处理加载、空态、错误、骨架四态；接口出错时返回友好提示而不暴露 Java 异常堆栈；Redis 未启动时核心 CRUD 功能仍可正常运行（代码中做降级处理）。")
h3("2.3.4  系统扩展性需求")
p("采用分层单体架构，便于横向扩展；预留 ProductDataSource / PriceDataSource 抽象接口，未来可接入真实电商数据源；算法权重集中配置、可动态调整。")
h3("2.3.5  系统安全性需求")
bullet("密码使用 BCrypt 哈希存储，禁止明文；")
bullet("采用 JWT 认证，JWT secret 不硬编码于业务代码；")
bullet("基于 RBAC 的前后端双重权限校验（普通用户 / 编辑 / 管理员），不能只在前端隐藏按钮；")
bullet("具备 SQL 注入防护、XSS 基础防护，敏感信息（密码、token）不对外返回；")
bullet("开发配置与生产配置分离，数据库密码等敏感配置使用环境变量。")

# ---------- 第3章 现状与设计条件 ----------
h1("第3章  现状与设计条件")
h2("3.1  行业与现状分析")
p("目前数码产品信息散落在电商平台、测评媒体与社区论坛，消费者需要跨多个渠道自行拼接信息。专业测评虽权威但篇幅长、术语多，普通消费者阅读门槛高；电商平台的比价功能多局限于单一平台内部，缺乏跨平台聚合与历史价格参考。因此，聚合「测评 + 口碑 + 多平台价格 + 历史价格」并以通俗方式输出的平台具有现实需求。")
h2("3.2  目标用户分析")
caption("表 3-1  目标用户画像表")
table(
    ["维度", "描述"],
    [
        ["年龄段", "18～35 岁，以学生和刚工作的年轻消费者为主"],
        ["技术水平", "会网购、会搜索，但看不懂 CPU/GPU/屏幕等专业参数的深层含义"],
        ["使用场景", "选购手机/笔记本/耳机等数码产品前，想快速知道「哪个值得买、现在买亏不亏」"],
        ["核心痛点", "多平台价格不一、专业参数难懂、不知当前是否好价、同价位产品难以抉择"],
    ],
    widths=[3.0, 10.0],
)
h2("3.3  技术条件")
p("本项目的技术选型遵循「成熟、开源、可演示、可降级」原则，具体选型及依据如表 3-2 所示。")
caption("表 3-2  技术选型依据表")
table(
    ["技术", "选型依据"],
    [
        ["Java 21 + Spring Boot 3.4", "主流企业级后端框架，生态成熟，满足毕业设计完整性"],
        ["MyBatis-Plus 3.5.9", "简化数据访问层，分页、逻辑删除等开箱即用"],
        ["MySQL 8.4", "关系型数据库，社区版免费，性能稳定"],
        ["Redis", "缓存加速，设计为可降级，不成为硬性依赖"],
        ["Vue 3 + TypeScript + Vite", "现代前端框架，类型安全，开发体验好"],
        ["Element Plus", "成熟的中后台组件库，快速搭建管理后台"],
        ["ECharts", "价格曲线、数据看板等图表的可视化方案"],
        ["JWT + Spring Security", "无状态认证与基于角色的访问控制"],
    ],
    widths=[5.0, 8.0],
)

# ---------- 第4章 总体设计方案 ----------
h1("第4章  总体设计方案")
h2("4.1  项目建设目标")
p("建设一个「高性价比数码产品测评聚合与比价系统」，以产品为中心，聚合测评、评价、价格三大类数据，通过性价比评分与个性化推荐两大算法，为普通消费者提供购买决策支持，同时为管理员提供完整的内容与数据管理能力。")
h2("4.2  设计思路与设计原则")
p("（1）前后端分离：前端（用户端、管理后台）与后端通过 RESTful API 通信，职责清晰、可独立部署；")
p("（2）分层架构：后端按 controller / service / mapper 分层，业务逻辑与算法集中在 service 层，便于维护与测试；")
p("（3）可降级设计：Redis 与 LLM 均为可选依赖，缺失时核心功能不受影响；")
p("（4）算法可解释：评分与推荐均基于明确的加权公式，便于论文论述与结果解释。")
h2("4.3  总体架构")
p("系统整体采用「前端展示层—后端服务层—数据存储层」的三层结构，如图 4-1 所示。")
caption("图 4-1  系统总体架构图")
bullet("展示层：用户端（Vue 3）、管理后台（Vue 3），通过 Axios 访问后端接口；")
bullet("服务层：Spring Boot REST API → Service（业务 + 算法）→ MyBatis-Plus → MySQL；Redis 提供缓存加速；")
bullet("数据层：MySQL 8 存储全部业务数据，初始化脚本包含表结构与基础/演示数据。")
caption("表 4-1  技术栈表")
table(
    ["层", "技术"],
    [
        ["后端", "Java 21、Spring Boot 3.4、MyBatis-Plus 3.5.9、MySQL 8.4、Redis（可降级）、JWT、Spring Security、Lombok、Hibernate Validator、Swagger"],
        ["用户端", "Vue 3、TypeScript、Vite、Element Plus、Pinia、Vue Router、Axios、ECharts"],
        ["管理后台", "Vue 3、TypeScript、Vite、Element Plus、Pinia、Vue Router、Axios、ECharts"],
        ["构建", "Maven（后端）、npm + Vite（前端）"],
    ],
    widths=[2.2, 10.8],
)
h2("4.4  系统总体布局")
caption("表 4-2  系统模块划分表")
table(
    ["子系统", "模块划分"],
    [
        ["用户端（frontend）", "首页、产品浏览/搜索、产品详情、比价、对比、测评、评价、收藏、浏览历史、排行榜、个人中心、推荐"],
        ["管理后台（admin）", "数据看板、用户/产品/分类/品牌/平台/价格/测评/评论管理、算法参数配置"],
        ["后端（backend）", "认证授权、产品、价格、内容、行为、算法、管理接口等 REST 服务"],
        ["数据库（database）", "24 张表，覆盖 RBAC、基础数据、产品、价格、内容、行为、推荐与算法配置"],
    ],
    widths=[3.6, 9.4],
)

# ---------- 第5章 业务数据流设计方案 ----------
h1("第5章  业务数据流设计方案")
h2("5.1  用户注册与登录流程")
p("用户在前端提交注册信息 → 后端校验用户名/邮箱唯一性 → 密码经 BCrypt 加密后写入 users 表 → 返回注册结果。登录时后端校验密码，成功后签发 JWT（含用户 ID 与角色），前端将 token 存入本地存储，后续请求携带 Authorization 头；token 过期或无效返回 401，前端跳转登录页。")
h2("5.2  产品搜索与比价流程")
p("用户输入关键词或选择筛选条件 → 前端组装查询参数 → 后端按条件（分类、品牌、价格区间、评分区间、排序）查询 products 表并关联品牌/分类 → 返回分页结果。进入产品详情后，后端查询该产品在 product_prices 表中各平台价格，计算全网最低价与价差，并返回 price_history 表中历史价格供 ECharts 绘图。")
h2("5.3  价格录入与历史记录流程")
p("管理员在后台为某产品某平台录入价格 → 后端写入 product_prices 表并更新 products 表的 current_price（全网最低价）→ 同时在 price_history 表插入一条历史记录，保证「更新价格即生成历史」的一致性。")
h2("5.4  评价与收藏流程")
p("登录用户对产品发表评价 → 后端校验内容与星级 → 写入 reviews 表并触发 products 表 user_rating 的聚合更新；用户收藏产品/测评时写入 favorites 表（唯一约束保证幂等）；浏览产品时后端自动写入 browse_history 表。上述行为同时作为推荐算法的输入信号。")
h2("5.5  性价比评分流程")
p("后端读取 algorithm_configs 表中的权重配置 → 计算产品在性能、功能、用户评分、价格优势、热度五个维度的得分（各归一化 0-100）→ 加权求和得到 0-100 的性价比评分 → 回写 products 表 value_score 字段。管理员修改权重后可触发全站重算。")
h2("5.6  个性化推荐流程")
p("后端聚合用户的浏览、收藏、评价行为形成偏好画像 → 对候选产品分别计算分类匹配、标签匹配、价格匹配、性价比、热度、新鲜度六个因素 → 归一化加权求和得到推荐分 → 写入 recommendations 表并返回推荐列表。无行为的新用户走冷启动，返回高性价比与热门产品。")

# ---------- 第6章 应用系统设计方案 ----------
h1("第6章  应用系统设计方案")
h2("6.1  用户系统模块")
p("提供注册、登录、退出、个人中心与修改密码。数据表：users、roles、user_roles。核心接口：POST /api/auth/register、POST /api/auth/login、GET /api/users/me。密码采用 BCrypt 存储，登录态采用 JWT。")
h2("6.2  产品浏览与搜索模块")
p("提供分类/品牌浏览、关键词搜索与多条件组合筛选、多方式排序。数据表：products、categories、brands、tags、product_tags。核心接口：GET /api/products（支持 keyword、categoryId、brandId、价格/评分区间、排序、分页）。")
h2("6.3  产品详情与多平台比价模块")
p("聚合单个产品的参数、评分、多平台价格与价差。数据表：products、product_specs、product_images、product_prices、platforms。核心接口：GET /api/products/{id}、GET /api/prices/compare/{id}。")
h2("6.4  历史价格分析模块")
p("展示产品在各平台的历史价格曲线与统计指标（当前最低、历史最低、历史最高、均价、降价幅度）。数据表：price_history。核心接口：GET /api/prices/history/{id}。前端使用 ECharts 折线图渲染。")
h2("6.5  产品横向对比模块")
p("用户选择 2～4 个产品进行参数、价格、评分并排对比。数据表：products、product_specs、comparison_records。核心接口：GET /api/products/compare?ids=1,2,3。")
h2("6.6  专业测评与用户评价模块")
p("测评由管理员/编辑发布，用户可查看、点赞、收藏、评论；评价由登录用户发表，支持点赞与评论/回复，允许一个用户对同一产品多次评价。数据表：review_articles、reviews、review_comments、review_likes、favorites。核心接口：GET /api/articles、GET /api/reviews/product/{id}、POST /api/reviews。")
h2("6.7  排行榜模块")
p("提供性价比榜、热门榜、降价榜、用户评分榜、最新产品榜，支持按分类查看，数据来自后端真实排序。核心接口：GET /api/products/ranking?type=value。")
h2("6.8  性价比评分算法")
p("性价比评分是毕业设计核心算法之一，综合性能、功能、用户评分、价格优势、热度五个维度，全部归一化到 0-100 后加权求和：")
p("性价比评分 = 性能得分×w1 + 功能得分×w2 + 用户评分×w3 + 价格优势×w4 + 热度×w5", indent=True)
p("其中：性能得分为产品性能指标；功能得分由产品参数条数映射（每 12 个参数满分）；用户评分由评价星级×20 得到；价格优势以 50 分为基准，越接近历史最低价、低于同类均价则加分。权重由管理后台配置，默认值如表 6-1 所示。")
caption("表 6-1  性价比评分权重默认配置表")
table(
    ["维度", "性能", "功能", "用户评分", "价格优势", "热度"],
    [
        ["默认权重", "30%", "20%", "20%", "30%", "0%"],
    ],
    widths=[3.0, 2.0, 2.0, 2.0, 2.0, 2.0],
)
p("采用多维加权而非简单的「性能/价格」比值，是因为单一比值忽略了用户口碑、功能丰富度与价格是否处于历史低位，多维加权能够解释每个维度的贡献，具备可解释性。")
h2("6.9  混合推荐算法")
p("推荐算法基于「用户行为 + 产品标签 + 价格区间 + 产品分类 + 性价比 + 热度 + 新鲜度」的混合策略：")
p("推荐分数 = 分类匹配×25 + 标签匹配×20 + 价格匹配×15 + 性价比×20 + 热度×10 + 新鲜度×10", indent=True)
caption("表 6-2  推荐算法各因素权重表")
table(
    ["因素", "分类匹配", "标签匹配", "价格匹配", "性价比", "热度", "新鲜度"],
    [
        ["权重", "25", "20", "15", "20", "10", "10"],
    ],
    widths=[3.0, 2.0, 2.0, 2.0, 2.0, 2.0, 2.0],
)
p("分类/标签匹配度采用 Jaccard 相似度计算；价格匹配度衡量候选价格与用户偏好均价的接近程度。无行为的新用户走冷启动，返回高性价比、热门、降价产品，保证推荐始终有结果。")
h2("6.10  管理后台")
p("管理后台采用左侧菜单 + 顶部导航 + 内容区 + 面包屑布局，覆盖用户、产品、分类、品牌、平台、价格、测评、评论的管理，以及算法参数配置与数据看板。所有管理接口位于 /api/admin/**，需 ADMIN 或 EDITOR 角色；普通用户访问返回 403。更新价格时自动生成历史价格记录。")

# ---------- 第7章 数据库设计 ----------
h1("第7章  信息资源规划与数据库设计方案")
h2("7.1  数据内容分析")
h3("7.1.1  数据需求")
p("系统数据可归纳为七类：权限数据（用户、角色、权限）、基础数据（分类、品牌、平台、标签）、产品数据（产品、参数、图片、标签关联）、价格数据（平台价格、历史价格）、内容数据（测评、评价、评论、点赞）、行为数据（收藏、浏览历史、对比记录）、算法数据（推荐结果、算法配置）。")
h3("7.1.2  数据分类")
p("上述七类数据中，产品数据是核心枢纽，价格数据与内容数据围绕产品展开，行为数据围绕用户展开，算法数据是评分与推荐的计算依据，权限数据为全系统提供访问控制。")
h2("7.2  数据库设计")
h3("7.2.1  数据设计标准")
bullet("数据库名 geek_price，字符集 utf8mb4，排序规则 utf8mb4_general_ci；")
bullet("存储引擎统一采用 InnoDB，支持事务与外键；")
bullet("主键统一为 BIGINT 自增 id；")
bullet("通用审计字段：created_at、updated_at；逻辑删除字段 deleted（0 未删 / 1 已删）；")
bullet("外键关系通过业务层维护，配合索引保证查询效率。")
h3("7.2.2  全局 ER 图")
p("系统包含 24 张表，其全局实体关系如图 7-1 所示（图中 1 表示「一」，N 表示「多」）。")
caption("图 7-1  极客比价系统全局实体关系图（ER 图）")
doc.add_picture(os.path.join(os.path.dirname(__file__), "er_diagram.png"),
               width=Cm(15.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
h3("7.2.3  数据库结构")
caption("表 7-1  数据库表分类清单表")
table(
    ["分类", "表名", "说明"],
    [
        ["RBAC 权限", "users", "用户表（用户名、密码、邮箱等）"],
        ["RBAC 权限", "roles", "角色表（USER/EDITOR/ADMIN）"],
        ["RBAC 权限", "permissions", "权限表"],
        ["RBAC 权限", "user_roles", "用户-角色关联表"],
        ["RBAC 权限", "role_permissions", "角色-权限关联表"],
        ["基础数据", "categories", "产品分类表（支持层级）"],
        ["基础数据", "brands", "品牌表"],
        ["基础数据", "platforms", "电商平台表"],
        ["基础数据", "tags", "产品标签表"],
        ["产品体系", "products", "产品表"],
        ["产品体系", "product_specs", "产品参数表"],
        ["产品体系", "product_images", "产品图片表"],
        ["产品体系", "product_tags", "产品-标签关联表"],
        ["价格体系", "product_prices", "产品平台价格表"],
        ["价格体系", "price_history", "价格历史表"],
        ["内容体系", "review_articles", "专业测评表"],
        ["内容体系", "reviews", "用户评价表"],
        ["内容体系", "review_comments", "评价评论表"],
        ["内容体系", "review_likes", "评价点赞表"],
        ["用户行为", "favorites", "收藏表（产品/测评）"],
        ["用户行为", "browse_history", "浏览历史表"],
        ["用户行为", "comparison_records", "对比记录表"],
        ["推荐与算法", "recommendations", "推荐结果表"],
        ["推荐与算法", "algorithm_configs", "算法参数配置表"],
    ],
    widths=[3.0, 4.0, 6.0],
)
h3("7.2.4  核心表结构")
p("以下列出六张核心业务表的关键字段结构。")
caption("表 7-2  users 用户表")
table(
    ["字段", "类型", "约束", "说明"],
    [
        ["id", "BIGINT", "PK", "主键"],
        ["username", "VARCHAR(50)", "NOT NULL, UK", "用户名（登录用）"],
        ["password", "VARCHAR(100)", "NOT NULL", "密码（BCrypt 哈希）"],
        ["nickname", "VARCHAR(50)", "-", "昵称"],
        ["email", "VARCHAR(100)", "UK", "邮箱"],
        ["phone", "VARCHAR(20)", "-", "手机号"],
        ["avatar", "VARCHAR(500)", "-", "头像 URL"],
        ["status", "TINYINT", "默认 1", "状态：1 正常 0 禁用"],
        ["created_at/updated_at", "DATETIME", "-", "创建/更新时间"],
        ["deleted", "TINYINT", "默认 0", "逻辑删除"],
    ],
    widths=[3.2, 3.0, 2.8, 4.0],
)
caption("表 7-3  products 产品表")
table(
    ["字段", "类型", "约束", "说明"],
    [
        ["id", "BIGINT", "PK", "主键"],
        ["name", "VARCHAR(200)", "NOT NULL", "产品名称"],
        ["brand_id / category_id", "BIGINT", "FK", "品牌 / 分类 ID"],
        ["official_price / current_price", "DECIMAL(12,2)", "-", "官方价 / 当前全网最低价"],
        ["score", "DECIMAL(5,1)", "-", "综合评分（0-10）"],
        ["value_score", "INT", "-", "性价比评分（0-100）"],
        ["performance_score / user_rating", "INT / DECIMAL", "-", "性能得分 / 用户评分"],
        ["sales / heat", "INT", "-", "销量 / 热度"],
        ["launch_date", "DATE", "-", "上市时间"],
        ["status / is_recommend / is_hot", "TINYINT", "-", "上架 / 推荐 / 热门标记"],
    ],
    widths=[4.0, 2.8, 2.4, 3.8],
)
caption("表 7-4  product_prices 产品平台价格表")
table(
    ["字段", "类型", "约束", "说明"],
    [
        ["id", "BIGINT", "PK", "主键"],
        ["product_id / platform_id", "BIGINT", "FK", "产品 / 平台 ID"],
        ["title / url", "VARCHAR", "-", "商品标题 / 链接"],
        ["price", "DECIMAL(12,2)", "NOT NULL", "当前价格"],
        ["original_price / discount", "DECIMAL(12,2)", "-", "原价 / 优惠金额"],
        ["shop_name", "VARCHAR(100)", "-", "店铺名称"],
        ["is_official", "TINYINT", "默认 0", "是否自营/官方"],
        ["stock_status", "TINYINT", "默认 1", "库存状态"],
    ],
    widths=[3.6, 2.8, 2.4, 4.2],
)
caption("表 7-5  price_history 价格历史表")
table(
    ["字段", "类型", "约束", "说明"],
    [
        ["id", "BIGINT", "PK", "主键"],
        ["product_id / platform_id", "BIGINT", "FK", "产品 / 平台 ID"],
        ["price", "DECIMAL(12,2)", "NOT NULL", "价格"],
        ["record_date", "DATE", "NOT NULL", "记录日期"],
        ["source", "VARCHAR(50)", "默认 manual", "数据来源"],
        ["created_at", "DATETIME", "-", "创建时间"],
    ],
    widths=[4.0, 3.0, 2.6, 3.4],
)
caption("表 7-6  reviews 用户评价表")
table(
    ["字段", "类型", "约束", "说明"],
    [
        ["id", "BIGINT", "PK", "主键"],
        ["product_id / user_id", "BIGINT", "FK", "产品 / 用户 ID"],
        ["star", "TINYINT", "默认 5", "星级（1-5）"],
        ["content", "VARCHAR(1000)", "-", "评价内容"],
        ["like_count", "INT", "默认 0", "点赞数"],
        ["created_at / updated_at", "DATETIME", "-", "创建/更新时间"],
    ],
    widths=[3.6, 3.0, 2.6, 3.8],
)
caption("表 7-7  review_articles 专业测评表")
table(
    ["字段", "类型", "约束", "说明"],
    [
        ["id", "BIGINT", "PK", "主键"],
        ["title / cover", "VARCHAR", "-", "标题 / 封面图"],
        ["product_id", "BIGINT", "FK", "关联产品 ID"],
        ["author / summary", "VARCHAR", "-", "作者 / 摘要"],
        ["content", "LONGTEXT", "-", "正文（富文本）"],
        ["score", "DECIMAL(4,1)", "-", "综合评分"],
        ["performance/appearance/screen/battery_score", "DECIMAL(4,1)", "-", "性能/外观/屏幕/续航评价"],
        ["pros / cons", "VARCHAR(1000)", "-", "优点 / 缺点"],
        ["views / likes", "INT", "-", "浏览量 / 点赞数"],
        ["status", "TINYINT", "默认 1", "状态：1 已发布 0 下架"],
    ],
    widths=[4.4, 2.6, 2.0, 4.0],
)

h2("7.3  数据库部署方案")
p("数据库部署于开发机本机 MySQL 8.4 实例，默认连接 localhost:3306，账号 root，密码 123456（可在 application.yml 中修改）。数据库通过三份 SQL 脚本完成初始化：01-schema.sql（表结构）、02-base-data.sql（基础数据）、03-product-data.sql（演示数据）。")
h2("7.4  数据量测算")
caption("表 7-8  数据量测算表（演示/初期规模）")
table(
    ["数据项", "规模", "估算依据"],
    [
        ["产品（products）", "132 条", "覆盖手机/笔记本/耳机等 16 类主要数码品类"],
        ["平台价格（product_prices）", "396 条", "每产品约 3 个平台"],
        ["历史价格（price_history）", "11,880 条", "每价格约 30 条历史记录"],
        ["专业测评（review_articles）", "32 篇", "重点产品覆盖"],
        ["用户评价（reviews）", "396 条", "每产品约 3 条"],
        ["用户（users）", "若干", "注册用户 + 默认账号"],
    ],
    widths=[4.6, 3.0, 5.4],
)

# ---------- 第8章 关键技术组件设计方案 ----------
h1("第8章  关键技术组件设计方案")
h2("8.1  安全认证（JWT + RBAC）")
p("系统采用「JWT + Spring Security + RBAC」的安全方案：登录成功后签发含用户 ID 与角色的 JWT，前端携带 Authorization 头访问接口；后端通过过滤器校验 token 并注入 SecurityContext；接口层通过角色注解控制访问权限。")
caption("表 8-1  角色权限表")
table(
    ["角色", "编码", "权限说明"],
    [
        ["普通用户", "USER", "浏览、搜索、评价、收藏、查看推荐等用户端功能"],
        ["编辑", "EDITOR", "在普通用户基础上，可管理测评内容"],
        ["管理员", "ADMIN", "全站数据与用户管理、算法参数配置"],
    ],
    widths=[3.0, 3.0, 7.0],
)
h2("8.2  缓存设计（Redis 可降级）")
p("Redis 用于缓存热点数据（如分类/品牌列表、产品详情、推荐结果）以加速访问。代码中对 Redis 不可用做了降级处理：连接失败或超时时自动回源数据库查询，因此 Redis 不成为系统硬性依赖，未启动时核心 CRUD 功能仍可正常运行。")
h2("8.3  AI 增强功能（LLM 可降级）")
p("系统预留三个轻量 LLM 增强能力：AI 产品摘要生成、AI 推荐理由解释、AI 对比总结。所有 AI 能力均设计为可降级——无 API key、断网、超时或额度耗尽时，AI 入口自动隐藏或置灰，核心功能不受影响，不依赖 LLM 服务。")
h2("8.4  前端可视化（ECharts）")
p("历史价格曲线、管理后台数据看板（用户增长趋势、产品分类分布、热门产品排行、价格变化趋势）均采用 ECharts 渲染，支持交互式图表的展示。")

# ---------- 第9章 系统部署方案 ----------
h1("第9章  系统部署方案")
h2("9.1  开发环境")
caption("表 9-1  开发环境配置表")
table(
    ["类别", "配置"],
    [
        ["操作系统", "Windows 11"],
        ["JDK", "Java 21"],
        ["构建工具", "Maven 3.9+、Node.js 22、npm"],
        ["数据库", "MySQL 8.4"],
        ["缓存", "Redis（可选）"],
        ["IDE", "VS Code / IntelliJ IDEA"],
    ],
    widths=[3.0, 10.0],
)
h2("9.2  部署架构")
p("开发/演示阶段，三个服务运行于同一台机器：后端 Spring Boot 运行于 8080 端口，用户端 Vite 开发服务器运行于 5173 端口，管理后台运行于 5174 端口，通过 Vite 代理将 /api 转发至后端。")
h2("9.3  系统软件配置")
p("生产部署时：后端通过 mvn package 打包为可执行 jar，以 nohup 方式运行；前端 npm run build 生成 dist 静态资源，交由 Nginx 托管，并由 Nginx 将 /api 反向代理至后端 8080 端口；数据库使用服务器 MySQL，连接信息通过 application-prod.yml 与环境变量配置。")

# ---------- 第10章 工程组织与运行管理 ----------
h1("第10章  工程组织与运行管理")
h2("10.1  项目组织")
p("本项目为个人毕业设计，由开发者一人承担需求分析、系统设计、编码实现、测试联调与文档撰写全部角色；技术路线与设计决策参考《总控 Prompt》与产品需求文档。")
h2("10.2  实施进度")
caption("表 10-1  项目实施进度计划表")
table(
    ["阶段", "主要工作", "交付物"],
    [
        ["需求分析", "明确目标用户、功能清单、优先级", "Product-Spec.md"],
        ["系统设计", "总体架构、数据库、算法设计", "本文档、schema.sql"],
        ["开发实现", "后端、用户端、管理后台编码", "backend/frontend/admin"],
        ["测试联调", "接口联调、功能验证、数据初始化", "可运行系统"],
        ["验收交付", "部署演示、文档整理、论文撰写", "演示环境 + 论文"],
    ],
    widths=[2.6, 6.4, 4.0],
)
h2("10.3  运行维护方案")
p("系统上线后，管理员可通过管理后台完成日常的内容更新（产品、价格、测评）与算法参数调整；价格更新自动生成历史记录，保证历史价格曲线持续有效；数据库通过定期备份保障数据安全。")
h2("10.4  测试账号")
caption("表 10-2  测试账号表")
table(
    ["角色", "用户名", "密码"],
    [
        ["管理员", "admin", "123456"],
        ["编辑", "editor", "123456"],
        ["普通用户", "user", "123456"],
    ],
    widths=[4.0, 4.0, 4.0],
)

# ---------- 第11章 项目风险与风险管理 ----------
h1("第11章  项目风险与风险管理")
h2("11.1  风险识别与分析")
caption("表 11-1  风险识别与对策表")
table(
    ["风险类别", "风险描述", "影响", "应对措施"],
    [
        ["技术风险", "真实电商价格数据源难以获取", "比价数据真实性受限", "第一阶段使用结构化模拟数据，预留数据源接口，未来接入真实数据源"],
        ["技术风险", "LLM API 不可用/超时/超配额", "AI 功能不可用", "AI 功能可降级，入口自动隐藏或置灰，不影响核心功能"],
        ["技术风险", "算法评分主观、缺乏基准", "评分可信度存疑", "采用多维归一化加权，保证公式可解释、权重可配置"],
        ["性能风险", "并发访问下的响应延迟", "用户体验下降", "Redis 缓存加速 + 数据库索引优化"],
        ["安全风险", "越权访问、数据泄露", "系统安全性受损", "BCrypt + JWT + RBAC 前后端双重校验"],
    ],
    widths=[2.2, 4.2, 2.6, 4.0],
)
h2("11.2  风险对策与管理")
p("针对上述风险，项目在设计阶段即确立「可降级、可解释、可配置」原则：外部依赖（数据源、LLM、Redis）一律可降级；算法一律公式化、权重可配置；权限一律前后端双重校验。通过上述措施将主要风险控制在可接受范围内。")

# ============================================================
# 第二篇 分阶段建设计划
# ============================================================
page_break()
para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = para.add_run("第二篇    分阶段建设计划")
set_eastasia(run._element, CN_HEAD); run.font.name = LATIN
run.font.size = Pt(18); run.font.bold = True
doc.add_paragraph()

h1("第12章  项目阶段划分")
h2("12.1  项目阶段划分")
p("项目按「需求—设计—开发—联调—验收」五个阶段划分，开发阶段进一步细分为 18 个开发子阶段（搭建工程、数据库、认证、产品、价格、内容、行为、算法、后台等），以保证开发有序推进、逐阶段可验证。")
caption("表 12-1  项目阶段划分表")
table(
    ["阶段", "阶段内容", "里程碑"],
    [
        ["一期：需求与设计", "需求分析、原型设计、总体/详细设计、数据库设计", "Product-Spec、本文档、schema.sql"],
        ["二期：核心开发", "后端分层搭建、认证授权、产品/价格/内容/行为/算法模块、用户端页面、管理后台", "可运行的前后端系统"],
        ["三期：联调验收", "接口联调、数据初始化、功能验证、部署演示、文档完善", "演示环境与毕业论文"],
    ],
    widths=[3.6, 6.4, 3.0],
)
h2("12.2  一期项目计划")
p("一期完成需求分析与系统设计：产出产品需求文档（含 13 项功能需求与 3 项 AI 增强需求）、UI 原型提示词、总体设计方案与 24 张表的数据库结构，为开发奠定基础。")
h2("12.3  二期项目说明")
p("二期完成核心开发与联调：按分层架构实现后端 REST 服务与两大核心算法（性价比评分、混合推荐），完成用户端与管理后台全部页面，初始化演示数据，实现三端完整联调。")

h1("第13章  测试与验收方案")
h2("13.1  测试范围与方法")
bullet("功能测试：覆盖用户系统、产品浏览/搜索、比价、历史价格、对比、测评、评价、收藏、排行榜、推荐、管理后台等全部模块；")
bullet("接口测试：通过 Swagger/curl 验证主要接口的请求响应与统一返回结构；")
bullet("权限测试：验证普通用户访问管理接口返回 403、未登录访问受限接口返回 401；")
bullet("异常测试：验证 404、空结果、参数校验失败等异常场景的友好提示；")
bullet("算法验证：验证性价比评分与推荐结果由真实公式计算、权重可配置且重算生效。")
h2("13.2  验收标准")
p("（1）用户端、管理后台、后端 API、数据库完整联调，核心功能全部可用；")
p("（2）性价比评分与个性化推荐两大算法真实运行、结果可解释；")
p("（3）权限控制、异常处理、可降级设计符合需求文档要求；")
p("（4）界面视觉达到现代化科技产品网站水平，支持 1280/1440/1920 分辨率正常显示。")

doc.save(os.path.join(os.path.dirname(__file__), "极客比价系统项目分析与设计方案.docx"))
print("DOCX saved:", os.path.join(os.path.dirname(__file__), "极客比价系统项目分析与设计方案.docx"))
